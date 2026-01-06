import sys
from docx import  Document
from pathlib import Path
import csv
from . import finance
import random
from Dialogs import dialogop



if getattr(sys, 'frozen', False):
    BASE_DIR = Path(sys.executable).parent
else:
    BASE_DIR = Path(__file__).parent.parent


ui_op_dialog = dialogop.Ui_Dialog()

DATA_DIR = BASE_DIR / 'datas'
VVINDOVVS1_DIR = DATA_DIR / 'vvindovvs1.csv'
VVINDOVVS2_DIR = DATA_DIR / 'vvindovvs2.csv'
VVINDOVS3_DIR = DATA_DIR / 'vvindovs3.csv'

def set_data_dir_do(new_base):
    global DATA_DIR, VVINDOVVS1_DIR, VVINDOVVS2_DIR, VVINDOVS3_DIR

    DATA_DIR = Path(new_base)
    VVINDOVVS1_DIR = DATA_DIR / 'vvindovvs1.csv' 
    VVINDOVVS2_DIR = DATA_DIR / 'vvindovvs2.csv'
    VVINDOVS3_DIR = DATA_DIR / 'vvindovs3.csv'
    


_theme = 'Mövzu'
_date = 'Tarix'
_usedtime = 'SərfOlunanZaman'
_note = 'Not'
_status = 'Status'


def write_vvindovvs1_progress(Date, Topic, Hours, Note):

    DATA_DIR.mkdir(exist_ok=True)
    if not VVINDOVVS1_DIR.exists():
        with open(VVINDOVVS1_DIR, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow([_date, _theme, _usedtime, _note ])

    with open(VVINDOVVS1_DIR, 'a', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow([Date, Topic, Hours, Note])


def write_math_progress(Date, Topic, Hours, Note, Status):

    DATA_DIR.mkdir(exist_ok=True)
    if not VVINDOVVS2_DIR.exists():
         with open(VVINDOVVS2_DIR, 'w', newline='', encoding='utf-8') as f:
             writer = csv.writer(f)
             writer.writerow([_date, _theme, _usedtime, _note, _status])

    with open(VVINDOVVS2_DIR, 'a', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow([Date, Topic, Hours, Note, Status])


def find_random_math_problem():

   
    with open(VVINDOVVS2_DIR, 'r', newline='', encoding='utf-8') as f:
        reader = csv.reader(f)
        next(reader)
        lists = list(reader)

        if not lists:
            return None, None
        
        note_done = [line for line in lists if line[4] != 'Done']

        if not note_done:
            return None, None
        
        choiced = random.choice(note_done)
        
        
    return choiced[1], choiced[3]

def done_topic_math(problem):
    lines = []

    with open(VVINDOVVS2_DIR, 'r', newline='', encoding='utf-8') as f:
        reader = csv.reader(f)
        header = next(reader)

        for line in reader:
            if line[3] == problem:
                line[4] = 'Done'
            lines.append(line)
   

    with open(VVINDOVVS2_DIR, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(lines)            


def export_word_file(name_of_page):
    
    doc = Document()

    doc.add_heading(f'{name_of_page} üzrə qeydlərim. ', level=1)
    doc.add_paragraph('')

    with open(VVINDOVVS1_DIR, 'r', newline='', encoding='utf-8') as f:
        reader = csv.DictReader(f)

        for line in reader:
            tema = line.get(_theme,'')
            date = line.get(_date, '') 
            note = line.get(_note, '')


            doc.add_heading(tema, level=2)
            doc.add_paragraph(date)
            doc.add_paragraph(note)

    return doc        








